"""Behavioral tests against real formats plus constrained network fixtures."""
import io
import base64
import zlib
import json
import unittest
import zipfile
from datetime import datetime
from unittest.mock import Mock, patch

from studio.ingestion import MAX_BYTES, fetch_sheet, inspect_source


def make_pdf():
    """An actual, text-based ruled table PDF (no PDF-writing dependency)."""
    stream = b"0.5 w\n50 650 m 350 650 l S\n50 620 m 350 620 l S\n50 590 m 350 590 l S\n50 560 m 350 560 l S\n50 560 m 50 650 l S\n200 560 m 200 650 l S\n350 560 m 350 650 l S\nBT /F1 12 Tf 60 630 Td (Region) Tj ET\nBT /F1 12 Tf 210 630 Td (Revenue) Tj ET\nBT /F1 12 Tf 60 600 Td (North) Tj ET\nBT /F1 12 Tf 210 600 Td (125) Tj ET\nBT /F1 12 Tf 60 570 Td (South) Tj ET\nBT /F1 12 Tf 210 570 Td (75) Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
    ]
    content = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(content))
        content += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(content)
    content += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    content += f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return content


def response(status=200, content=b"Region,Revenue\nNorth,120\n", **headers):
    result = Mock()
    result.status_code = status
    result.headers = {"Content-Type": "text/csv", **headers}
    result.iter_content.return_value = iter([content])
    return result


class CSVImportTests(unittest.TestCase):
    def test_csv_quoted_fields_missing_and_duplicate_headers_are_retained(self):
        result = inspect_source("sales.csv", b'Region,Revenue,Revenue,Note\n"North, East",100,20,"line\nbreak"\nSouth,,5,\n')
        table = result["tables"][0]
        self.assertEqual(table["columns"], ["Region", "Revenue", "Revenue (2)", "Note"])
        self.assertEqual(table["rows"][0]["Region"], "North, East")
        self.assertEqual(table["rows"][0]["Note"], "line\nbreak")
        self.assertIsNone(table["rows"][1]["Revenue"])
        self.assertTrue(any("repeated" in warning for warning in result["warnings"]))

    def test_semicolon_and_windows_encoding(self):
        data = "City;Revenue\nMontréal;200\nZürich;150\n".encode("cp1252")
        result = inspect_source("sales.csv", data)
        self.assertEqual(result["tables"][0]["rows"][0], {"City": "Montréal", "Revenue": "200"})
        self.assertTrue(any("1252" in warning for warning in result["warnings"]))

    def test_utf16_and_tab_detection(self):
        result = inspect_source("sales.tsv", "Region\tRevenue\nNorth\t20\n".encode("utf-16"))
        self.assertEqual(result["tables"][0]["rows"][0]["Revenue"], "20")

    def test_extra_columns_and_short_rows_preserved(self):
        result = inspect_source("sales.csv", b"Region,Revenue\nNorth,10,bonus\nSouth\n")
        table = result["tables"][0]
        self.assertEqual(table["columns"], ["Region", "Revenue", "Column 3"])
        self.assertEqual(table["rows"][0]["Column 3"], "bonus")
        self.assertIsNone(table["rows"][1]["Revenue"])
        self.assertTrue(any("Inconsistent" in warning for warning in result["warnings"]))

    def test_header_override_and_skipped_rows(self):
        result = inspect_source("sales.csv", b"Annual report\nmetadata\nRegion,Revenue\nNorth,100\n", {"skip_rows": 1, "header_row": 1})
        self.assertEqual(result["tables"][0]["rows"], [{"Region": "North", "Revenue": "100"}])
        self.assertEqual(result["tables"][0]["header_row"], 2)

    def test_automatic_report_header_detection(self):
        table = inspect_source("sales.csv", b"Annual report,,\nRegion,Revenue,Cost\nNorth,100,20\nSouth,80,30\n")["tables"][0]
        self.assertEqual(table["columns"], ["Region", "Revenue", "Cost"])
        self.assertEqual(len(table["rows"]), 2)

    def test_malformed_quoting_is_actionable(self):
        with self.assertRaisesRegex(ValueError, "malformed quoting"):
            inspect_source("bad.csv", b'Name,Value\n"unterminated,100\n')

    def test_explicit_delimiter(self):
        table = inspect_source("pipe.csv", b"a|b\nx|1\n", {"delimiter": "|"})["tables"][0]
        self.assertEqual(table["columns"], ["a", "b"])


class ExcelImportTests(unittest.TestCase):
    def test_real_workbook_multiple_sheets_dates_booleans_and_empty_sheet(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sales"
        sheet.append(["Annual revenue"])
        sheet.append(["Region", "Revenue", "Date", "Active"])
        sheet.append(["North", 125.5, datetime(2025, 1, 2), True])
        sheet.append(["South", 75, datetime(2025, 1, 3), False])
        second = workbook.create_sheet("Targets")
        second.append(["Region", "Target"])
        second.append(["North", 200])
        workbook.create_sheet("Empty")
        output = io.BytesIO()
        workbook.save(output)
        result = inspect_source("sales.xlsx", output.getvalue())
        self.assertEqual([table["name"] for table in result["tables"]], ["Sales", "Targets"])
        first = result["tables"][0]
        self.assertEqual(first["key"], "sheet-0")
        self.assertEqual(first["rows"][0]["Revenue"], 125.5)
        self.assertEqual(first["rows"][0]["Date"], "2025-01-02T00:00:00")
        self.assertTrue(first["rows"][0]["Active"])
        json.dumps(result, allow_nan=False)

    def test_workbook_table_keys_stable_after_override(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["a", "b"])
        workbook.active.append([1, 2])
        output = io.BytesIO()
        workbook.save(output)
        default = inspect_source("data.xlsx", output.getvalue())
        custom = inspect_source("data.xlsx", output.getvalue(), {"header_row": 0})
        self.assertEqual(default["tables"][0]["key"], custom["tables"][0]["key"])

    def test_actual_legacy_xls_multiple_sheets(self):
        # Binary BIFF8/OLE fixture created once with xlwt 1.3.0. Reading the
        # fixture needs only the application's xlrd dependency.
        content = zlib.decompress(base64.b64decode("eJztWE1IVFEU/u6dGWdG/BlNAxVkELIyDaJNG5uMKTcppkE/BDXqQ0WZieco2KI0m2UQtCraCLZoY7XphwoKWhQERi2CIHBq2SooaKG+zjnvjY42grPQUN433HPvPeeec++759zz7rwPs2VzU4+q01iFw/Bg0QqiIIunqAQznRBIblnczNQBKpaLLYVggBxZ4MPz4vd+9iH7Ow2Nh97XRIFvVM7jEtoTcSO8iTgqa4gpXkOz8lLsadylUoIqWVe50B6hO4Q+kLEvhB4Rzg2hzTR2Tp3DbKS94ZATx2d0nchKiCo8EZ0vwjmASrzlOL56U9ljfWgxB2JD20VQ6y3CNMjXrUbcMLlfinsoBM4SmtramqLROVRQHEzjtxUGfmUO/Kuwy99cvgLx/6zk+9fgB8iJufjBHPxb2guMw7qoOPpTNGrMy5IinDD6Yj1j4a7YkDGcQikmJP37cSpm9hnJ4XlE5QXAmYIknUbfQCLu58aoER8xyEg0ljRI0tKTHBg1fKC0YSb7qe5KjCT7SWAbomiDnXhCKxJPsRzHIqK9NDm3y+RQhujx5u///NjW3RG5IJxxednYr6Rd/DiwMMEapFwiEo9QljSIxj6h18RqjbSrhVbQwaC6vqPSaRyflDHXRVpP8xwUfIrszmrvoXbqx8mntanvkb3UnmlNX66Y+RyZQh3tUC/p828SjapR3bnNeBbJ1MpJXl+FVv2TyAI65Kzdsj2mSrHAJ5RQJtTu8e6opZ529srueajncWypHLaU2OLdeUmrUWKr5g2n1fg7TXZ8OkA8j8zANnQOG1ps+Bwbtrf2K34Kc8mGFht0ndDleMwKlJqXsRlxoJfoVo0DvY44KMjy9erxy76+ogsdX9d5HJ8gl09cuHDhwoULF9sa9oXBvi7zFYGvc3yZ4LsYf9dZoLL4vz9SuNgwdCJBvyT9QzyGONUmxvKKn53wqYwttU6dzPdCxmma3cQgumUdg3nNzaArtMp+nnUrhvKeai3kPf9iPuvc4Pn/AsVoH6I="))
        result = inspect_source("legacy.xls", content)
        self.assertEqual([table["name"] for table in result["tables"]], ["Legacy Sales", "Targets"])
        self.assertEqual(result["tables"][0]["rows"][0]["Revenue"], 125.5)
        self.assertEqual(result["tables"][0]["rows"][0]["Date"], "2025-01-02T00:00:00")
        self.assertTrue(result["tables"][0]["rows"][0]["Active"])

    def test_invalid_xls_and_corrupted_xlsx(self):
        with self.assertRaisesRegex(ValueError, "valid legacy Excel"):
            inspect_source("bad.xls", b"a,b\n1,2")
        with self.assertRaisesRegex(ValueError, "corrupted"):
            inspect_source("bad.xlsx", b"not a zip")

    def test_zip_bomb_rejected_before_workbook_parser(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", "a" * (2 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "unsafe compressed size"):
            inspect_source("bomb.xlsx", output.getvalue())


class JSONImportTests(unittest.TestCase):
    def test_nested_objects_arrays_and_api_metadata_preserved(self):
        content = {"meta": {"total": 2}, "data": [{"customer": {"name": "A"}, "amount": 10, "items": [{"sku": "x", "qty": 2}, {"sku": "y", "qty": 1}]}, {"customer": {"name": "B"}, "amount": 20, "items": [{"sku": "z", "qty": 3}]}]}
        result = inspect_source("api.json", json.dumps(content).encode())
        tables = {table["key"]: table for table in result["tables"]}
        self.assertIn("$", tables)
        self.assertIn("$.data", tables)
        self.assertIn("$.data[].items", tables)
        self.assertEqual(tables["$"]["rows"][0]["meta.total"], 2)
        self.assertEqual(tables["$.data"]["rows"][0]["customer.name"], "A")
        self.assertEqual(json.loads(tables["$.data"]["rows"][0]["items"])[0]["sku"], "x")
        self.assertEqual(tables["$.data[].items"]["rows"][2]["__parent_index"], 1)
        self.assertEqual(len(tables["$.data[].items"]["rows"]), 3)

    def test_mixed_fields_union_and_scalar_arrays(self):
        tables = inspect_source("data.json", b'[{"a":1},{"b":true},null]')["tables"]
        self.assertEqual(tables[0]["columns"], ["a", "b", "value"])
        self.assertEqual(tables[0]["rows"][1], {"a": None, "b": True, "value": None})
        scalar = inspect_source("values.json", b'[10,20,30]')["tables"][0]
        self.assertEqual([row["value"] for row in scalar["rows"]], [10, 20, 30])

    def test_invalid_json_duplicate_keys_and_nonfinite_values_rejected(self):
        for value in [b'{"a":', b'{"a":1,"a":2}', b'[NaN]', b'[Infinity]']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                inspect_source("bad.json", value)

    def test_large_api_response_array_remains_a_selectable_table(self):
        records = [{"category": "North region", "revenue": i} for i in range(4000)]
        result = inspect_source("api.json", json.dumps({"meta": {"total": 4000}, "data": records}).encode())
        tables = {table["key"]: table for table in result["tables"]}
        self.assertEqual(len(tables["$.data"]["rows"]), 4000)
        self.assertEqual(tables["$.data"]["rows"][-1]["revenue"], 3999)
        self.assertIn("Nested table", tables["$"]["rows"][0]["data"])
        self.assertTrue(any("references when oversized" in warning for warning in result["warnings"]))

    def test_overly_nested_json_is_rejected(self):
        value = {"field": 1}
        for _ in range(40):
            value = {"child": value}
        with self.assertRaisesRegex(ValueError, "nesting"):
            inspect_source("deep.json", json.dumps(value).encode())

    def test_literal_dotted_key_collision_rejected(self):
        with self.assertRaisesRegex(ValueError, "collide"):
            inspect_source("data.json", b'[{"a.b":1,"a":{"b":2}}]')


class PDFImportTests(unittest.TestCase):
    def test_real_pdf_table_extraction(self):
        result = inspect_source("sales.pdf", make_pdf())
        table = result["tables"][0]
        self.assertEqual(table["columns"], ["Region", "Revenue"])
        self.assertEqual(table["rows"], [{"Region": "North", "Revenue": "125"}, {"Region": "South", "Revenue": "75"}])
        self.assertEqual(table["key"], "page-1-table-1")
        self.assertTrue(any("heuristic" in warning for warning in result["warnings"]))

    def test_no_detected_tables_explicitly_reports_ocr_limitation(self):
        fake_pdf = Mock()
        fake_pdf.pages = [Mock()]
        fake_pdf.pages[0].extract_tables.return_value = []
        fake_pdf.pages[0].extract_text.return_value = ""
        with patch("pdfplumber.open") as opening:
            opening.return_value.__enter__.return_value = fake_pdf
            with self.assertRaisesRegex(ValueError, "OCR"):
                inspect_source("scan.pdf", b"%PDF-1.4\nimage content")

    def test_invalid_pdf(self):
        with self.assertRaisesRegex(ValueError, "valid PDF"):
            inspect_source("bad.pdf", b"this is not a pdf")


class GoogleSheetTests(unittest.TestCase):
    URL = "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuv/edit#gid=123"

    @patch("studio.ingestion.requests.get")
    def test_sheet_id_and_gid_are_reconstructed_not_forwarded(self, get):
        get.return_value = response()
        name, content = fetch_sheet(self.URL)
        self.assertEqual(name, "google-sheet.csv")
        self.assertIn(b"North,120", content)
        called = get.call_args
        self.assertEqual(called.args[0], "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuv/export?format=csv&gid=123")
        self.assertFalse(called.kwargs["allow_redirects"])
        self.assertNotIn("verify", called.kwargs)  # requests trust verification stays enabled
        self.assertTrue(get.return_value.close.called)

    @patch("studio.ingestion.requests.get")
    def test_untrusted_urls_are_rejected_without_network_request(self, get):
        for url in ["http://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuv/edit", "https://evil.test/spreadsheets/d/abcdefghijklmnopqrstuv/edit", "https://docs.google.com.evil.test/spreadsheets/d/abcdefghijklmnopqrstuv/edit", "https://user:pass@docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuv/edit", "https://docs.google.com:8080/spreadsheets/d/abcdefghijklmnopqrstuv/edit", "https://127.0.0.1/", self.URL.replace("gid=123", "gid=bad")]:
            with self.subTest(url=url), self.assertRaises(ValueError):
                fetch_sheet(url)
        get.assert_not_called()

    @patch("studio.ingestion.requests.get")
    def test_private_permission_failure_is_actionable(self, get):
        get.return_value = response(status=403)
        with self.assertRaisesRegex(ValueError, "private"):
            fetch_sheet(self.URL)

    @patch("studio.ingestion.requests.get")
    def test_private_login_redirect_is_not_followed(self, get):
        get.return_value = response(status=302, Location="https://accounts.google.com/login")
        with self.assertRaisesRegex(ValueError, "requires permission"):
            fetch_sheet(self.URL)
        self.assertEqual(get.call_count, 1)

    @patch.dict("os.environ", {"GOOGLE_SHEETS_ACCESS_TOKEN": "test-only-server-token"})
    @patch("studio.ingestion.requests.get")
    def test_token_not_forwarded_to_download_host(self, get):
        get.side_effect = [response(status=302, Location="https://doc-12-ab.sheets.googleusercontent.com/export/data.csv"), response()]
        fetch_sheet(self.URL)
        self.assertIn("Authorization", get.call_args_list[0].kwargs["headers"])
        self.assertNotIn("Authorization", get.call_args_list[1].kwargs["headers"])

    @patch("studio.ingestion.requests.get")
    def test_ssrf_redirect_is_blocked(self, get):
        for location in ["http://127.0.0.1/private", "https://evil.test/export", "https://sheets.googleusercontent.com.evil.test/export", "https://foo.sheets.googleusercontent.com:8443/export"]:
            get.reset_mock()
            get.return_value = response(status=302, Location=location)
            with self.subTest(location=location), self.assertRaises(ValueError):
                fetch_sheet(self.URL)
            self.assertEqual(get.call_count, 1)

    @patch("studio.ingestion.requests.get")
    def test_oversized_download_rejected(self, get):
        get.return_value = response(**{"Content-Length": str(MAX_BYTES + 1)})
        with self.assertRaisesRegex(ValueError, "25 MB"):
            fetch_sheet(self.URL)

    @patch("studio.ingestion.requests.get")
    def test_html_login_response_rejected(self, get):
        get.return_value = response(content=b"<!DOCTYPE html><html>Sign in</html>")
        with self.assertRaisesRegex(ValueError, "sign-in"):
            fetch_sheet(self.URL)


class ImportBoundTests(unittest.TestCase):
    def test_empty_unsupported_and_invalid_options(self):
        for filename, data, options in [("data.csv", b"", None), ("data.exe", b"hello", None), ("data.csv", b"a,b\n1,2", {"skip_rows": -1}), ("data.csv", b"a,b\n1,2", {"header_row": True}), ("data.csv", b"a,b\n1,2", {"delimiter": "xx"})]:
            with self.subTest(filename=filename, options=options), self.assertRaises(ValueError):
                inspect_source(filename, data, options)

    def test_row_limit_rejects_instead_of_silently_truncating(self):
        with patch("studio.ingestion.MAX_ROWS", 2):
            with self.assertRaisesRegex(ValueError, "exceeds"):
                inspect_source("data.csv", b"Name,Amount\nA,1\nB,2\nC,3\n")

    def test_cell_and_upload_bounds(self):
        with patch("studio.ingestion.MAX_CELL_CHARS", 4):
            with self.assertRaises(ValueError):
                inspect_source("data.csv", b"Name,Value\nTooLong,1\n")
        with patch("studio.ingestion.MAX_BYTES", 3):
            with self.assertRaisesRegex(ValueError, "25 MB"):
                inspect_source("data.csv", b"a\nb\n")
